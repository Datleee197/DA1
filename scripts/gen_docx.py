import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    return h

def add_paragraph(text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    return p

def add_bullet(text):
    doc.add_paragraph(text, style='List Bullet')

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("CHƯƠNG 6. TRIỂN KHAI RTL VÀ KIỂM CHỨC NĂNG HỆ THỐNG FFT-256 R2²SDF")
run.bold = True
run.font.size = Pt(16)

add_heading('6.1. Mục tiêu của chương', 2)
add_paragraph('Chương 5 đã trình bày các kiến thức nền tảng và nguyên lý hoạt động của kiến trúc R2²SDF (Radix-2² Single-path Delay Feedback) ở mức độ thuật toán. Tiếp nối nền tảng đó, Chương 6 sẽ trình bày quá trình hiện thực hóa (triển khai) kiến trúc đó thành các mô-đun phần cứng (RTL - Register Transfer Level) cụ thể sử dụng ngôn ngữ SystemVerilog.')
add_paragraph('Chương này trình bày tiến trình thiết kế theo phương pháp từ dưới lên (Bottom-Up), bắt đầu từ các khối tính toán cơ sở (butterfly, rotator, multiplier), tiến tới lắp ghép thành các tầng SDF, đóng gói thành các khối R2²SDF hoàn chỉnh, và cuối cùng là liên kết tạo thành toàn bộ hệ thống FFT-16 và FFT-256.')
add_paragraph('Mục tiêu cốt lõi của chương này là cung cấp minh chứng rõ ràng rằng thiết kế không chỉ khả thi trên sơ đồ nguyên lý mà đã được kiểm chứng chức năng bằng mô phỏng RTL (RTL simulation) so với mô hình tham chiếu lý tưởng (golden model) được xây dựng bằng Python/NumPy.')

add_heading('6.2. Quy ước thiết kế RTL', 2)
add_paragraph('Toàn bộ dự án tuân theo các quy ước thiết kế đồng nhất để đảm bảo khả năng mở rộng và dễ dàng gỡ lỗi:')
add_bullet('Cấu trúc Giao tiếp (Interface): Tất cả các mô-đun đều sử dụng xung nhịp đồng bộ clk, tín hiệu reset tích cực mức thấp rst_n. Luồng dữ liệu được điều phối bởi các tín hiệu bắt tay (handshake) gồm valid_in (báo hiệu có dữ liệu hợp lệ) và sync_in (đánh dấu mẫu đầu tiên của một khung dữ liệu - frame). Ngõ ra tương ứng có valid_out và sync_out.')
add_bullet('Định dạng Dữ liệu: Mạch sử dụng số thực dấu phẩy tĩnh (fixed-point) định dạng Q1.15 có dấu (signed 16-bit). Một mẫu số phức (complex sample) bao gồm 16 bit phần thực và 16 bit phần ảo. Không sử dụng kiểu dữ liệu số thực dấu phẩy động (floating-point) để tối ưu hóa tài nguyên phần cứng.')
add_bullet('Giao thức luồng dữ liệu (Streaming Protocol): Để hệ thống đường ống (pipeline) hoạt động và dịch chuyển dữ liệu, tín hiệu valid_in bắt buộc phải được giữ ở mức cao. Nếu valid_in xuống mức thấp, toàn bộ pipeline sẽ bảo toàn trạng thái (stall).')
add_bullet('Giới hạn (Saturation): Để tiết kiệm logic, bản Minimum Viable Product (MVP) này không sử dụng logic bão hòa (saturation). Quá trình giảm tỷ lệ (scaling) liên tục được áp dụng để chủ động ngăn chặn tràn số.')

add_heading('6.3. Thiết kế các block RTL cơ bản', 2)

add_heading('6.3.1. trivial_rotator.sv', 3)
add_paragraph('Đây là khối thực hiện phép nhân số phức với các hệ số xoay đặc biệt mang tính chất trực giao (W^0, W^(N/4), W^(N/2), W^(3N/4)), tương đương với việc nhân cho 1, -j, -1, và j.')
add_paragraph('Thay vì sử dụng bộ nhân phần cứng đắt đỏ, khối này hoạt động hoàn toàn bằng tổ hợp (combinational logic) thông qua việc hoán đổi vị trí phần thực/phần ảo và đảo dấu (bù 2). Mối quan hệ giữa tín hiệu điều khiển rot_sel và phép toán như sau:')
table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Tín hiệu rot_sel'
hdr_cells[1].text = 'Phép toán tương đương'
row = table.add_row().cells
row[0].text = '0'
row[1].text = 'Nhân với 1'
row = table.add_row().cells
row[0].text = '1'
row[1].text = 'Nhân với -j'
row = table.add_row().cells
row[0].text = '2'
row[1].text = 'Nhân với -1'
row = table.add_row().cells
row[0].text = '3'
row[1].text = 'Nhân với j'
add_paragraph('Khối này đóng vai trò quyết định trong tầng BF2II của kiến trúc R2²SDF, cho phép xử lý phép xoay -j tự động mà không cần can thiệp từ bộ nhân phức bên ngoài.')

add_heading('6.3.2. butterfly_radix2_scaled.sv', 3)
add_paragraph('Khối thực hiện phép tính bướm Radix-2 chuẩn cho cả phần thực và phần ảo. Tuy nhiên, điểm khác biệt là nó tự động điều chỉnh biên độ (scaling):')
add_bullet('y0 = (a + b) >>> 1')
add_bullet('y1 = (a - b) >>> 1')
add_paragraph('Trước khi cộng hoặc trừ, các toán hạng 16-bit được mở rộng dấu (sign-extend) lên 17-bit. Sau đó kết quả được dịch phải (Arithmetic Right Shift) 1 bit. Quá trình này giúp ngăn chặn sự tăng trưởng bit (bit growth) dọc theo chuỗi tính toán kéo dài, nhưng đồng thời tạo ra hệ số tỷ lệ (scaling factor) tổng quát là 1/2 cho mỗi tầng bướm.')

add_heading('6.3.3. complex_multiplier_q15.sv', 3)
add_paragraph('Khối này thực hiện phép nhân của hai số phức định dạng Q1.15. Phương trình tính toán:')
add_bullet('y_re = z_re * tw_re - z_im * tw_im')
add_bullet('y_im = z_re * tw_im + z_im * tw_re')
add_paragraph('Phép nhân hai số 16-bit (Q1.15) sẽ tạo ra kết quả 32-bit (Q2.30). Mạch tiến hành cắt và trích xuất dải bit phù hợp để dịch ngược (shift right) về chuẩn Q1.15. Để tối ưu tần số hoạt động (timing), ngõ ra của bộ nhân được chốt qua một tầng thanh ghi (1-cycle registered output), làm tăng tổng độ trễ của datapath thêm 1 chu kỳ cho mỗi khối CMUL.')

add_heading('6.3.4. delay_line_shift.sv', 3)
add_paragraph('Đây là một thanh ghi dịch (shift register) tuyến tính tổng quát với ngõ ra được chốt bằng thanh ghi (registered output). Khối này được sử dụng để trì hoãn các tín hiệu điều khiển (như valid_in, sync_in) chạy song song với luồng dữ liệu chính.')

add_heading('6.3.5. sdf_feedback_delay.sv', 3)
add_paragraph('Đây là thanh ghi trễ chuyên dụng dùng cho luồng phản hồi (feedback) trong các tầng SDF. Khác với delay_line_shift, khối này sử dụng ngõ ra tổ hợp (combinational output) xuất trực tiếp từ đầu thanh ghi dịch. Thiết kế này giúp giải quyết hiện tượng trượt pha biên (phase-boundary skew) phát hiện được trong quá trình kiểm chứng, đảm bảo dữ liệu phản hồi khớp chính xác với luồng vào ở chu kỳ tiếp theo.')

add_heading('6.3.6. twiddle_rom.sv', 3)
add_paragraph('Khối bộ nhớ ROM tĩnh lưu trữ các hệ số lượng giác (twiddle factors) dưới định dạng Q1.15. Các giá trị này được tổng hợp từ các file mã thập lục phân (hex files) như tw256.hex, tw64.hex. Các file hex này được sinh ra tự động bởi kịch bản scripts/gen_twiddle_hex.py, giúp RTL không cần phải tính toán hàm sin/cos hoặc sử dụng số thực dấu phẩy động phức tạp.')

add_heading('6.4. Thiết kế stage SDF', 2)

add_heading('6.4.1. sdf_bf2i_stage.sv', 3)
add_paragraph('Tầng BF2I đảm nhiệm khoảng cách chéo $N/2$. Nó bao gồm bộ bướm Radix-2 kết hợp cùng bộ trễ phản hồi sdf_feedback_delay. Hoạt động dựa trên tín hiệu phase_sel:')
add_bullet('Pha Nạp (Fill phase): Dữ liệu đầu vào din được ghi vào bộ trễ. Ngõ ra của tầng chính là dữ liệu cũ bị xả ra từ bộ trễ (delay_out).')
add_bullet('Pha Tính Toán (Compute phase): Bộ bướm nhận ngõ vào a là dữ liệu từ bộ trễ, và b là dữ liệu đầu vào hiện tại (din). Nhánh y0 được đẩy thẳng ra ngõ ra, trong khi nhánh y1 được phản hồi ngược vào lại bộ trễ.')
add_paragraph('Giai đoạn kiểm chứng khối này đã phát hiện một rủi ro về độ đồng bộ tín hiệu: thanh ghi dịch của cờ valid/sync bắt buộc phải có độ sâu là DELAY (chứ không phải DELAY+1) để bù trừ chính xác với thời gian dữ liệu thực sự thoát ra khỏi bộ đệm tổ hợp.')

add_heading('6.4.2. sdf_bf2ii_stage.sv', 3)
add_paragraph('Tương tự như BF2I, tầng BF2II xử lý khoảng cách chéo $N/4$. Điểm làm nên lợi thế đặc trưng của kiến trúc Radix-2² là sự có mặt của khối trivial_rotator được đặt ngay ngõ vào bướm của tầng BF2II. Khối này sẽ tính toán phép xoay -j (nếu cần) thông qua tín hiệu rot_sel trước khi đưa vào bướm, giúp tầng BF2II thực hiện phép tính tương đương của chặng sau trong thuật toán Radix-4.')

add_heading('6.5. Thiết kế block Radix-2² SDF', 2)

add_heading('6.5.1. r22sdf_block.sv', 3)
add_paragraph('Đây là khối cấu trúc liên kết trực tiếp datapath theo trật tự: BF2I -> BF2II -> CMUL (tuỳ chọn). Biến HAS_CMUL=1 kích hoạt bộ nhân phức, trong khi HAS_CMUL=0 cho phép luồng dữ liệu truyền qua trực tiếp. Bản thân khối này không tự sinh tín hiệu điều khiển, mọi tín hiệu phase_sel_i, phase_sel_ii, rot_sel và tw_addr đều được cấp từ bên ngoài.')
add_paragraph('[Hình 6.1: Placeholder - Sơ đồ block R2²SDF cấu tạo từ BF2I, BF2II và CMUL]')

add_heading('6.5.2. r22sdf_block_controlled.sv', 3)
add_paragraph('Khối bọc (wrapper) này nâng cấp r22sdf_block bằng cách tích hợp logic điều khiển cục bộ (local distributed control). Thay vì dùng chung một bộ đếm toàn cục (global counter) cho hệ thống có độ trễ lên đến hàng trăm chu kỳ - rất dễ sinh lỗi định tuyến timing - mỗi khối đều có một bộ đếm current_cnt chạy song song với luồng dữ liệu valid_in/sync_in đi qua nó.')
add_paragraph('Các tín hiệu được ánh xạ:')
add_bullet('phase_sel_i và phase_sel_ii được trích xuất từ các bit tương ứng với cấp độ sâu của bộ trễ trong current_cnt.')
add_bullet('rot_sel được phái sinh logic trực tiếp từ phase_sel_i.')
add_bullet('tw_addr được sinh ra từ chỉ số cmul_idx, bộ chỉ số này được trễ hóa một khoảng thời gian bằng tổng độ trễ của tầng BF2I và BF2II để đảm bảo địa chỉ ROM được gọi đúng lúc dữ liệu đi tới khối CMUL.')
add_paragraph('Cơ chế này giữ nguyên bản chất đường đi dữ liệu như trên lý thuyết (Figure 4) nhưng làm cho hoạt động định thời (timing) của RTL trở nên mạnh mẽ và chắc chắn hơn.')

add_heading('6.6. Thiết kế flow từ block nhỏ lên hệ thống', 2)
add_paragraph('Phương pháp thiết kế Bottom-Up (từ dưới lên) được áp dụng một cách nghiêm ngặt. Việc thiết kế và kiểm chứng được tiến hành tuần tự qua 10 bước:')
add_paragraph('1. trivial_rotator\n2. butterfly_radix2_scaled\n3. complex_multiplier_q15\n4. Các module trễ (delay modules)\n5. Các tầng BF2I, BF2II\n6. Khối r22sdf_block\n7. Khối điều khiển r22sdf_block_controlled\n8. Môi trường sandbox FFT-16\n9. Hệ thống top-level FFT-256\n10. Đối chiếu chức năng hệ thống với mô hình golden Python.')
add_paragraph('Lợi ích của quy trình này là khoanh vùng và triệt tiêu lỗi (như lỗi trễ pha, lỗi tràn số) ngay tại cấp độ mô-đun con, đảm bảo việc tích hợp lên cấp độ hệ thống diễn ra trơn tru mà không gặp phải rủi ro sửa lỗi dạng "dây chuyền".')
add_paragraph('[Hình 6.2: Placeholder - Flow kiểm chứng bottom-up]')

add_heading('6.7. Thiết kế top-level FFT-16 sandbox', 2)
add_paragraph('Hệ thống FFT-16 được phát triển như một môi trường thử nghiệm trung gian (sandbox) nhằm đánh giá khả năng móc xích nhiều khối R2²SDF. Nó gồm 2 khối:')
add_bullet('Khối 0: DELAY_I=8, DELAY_II=4, HAS_CMUL=1, ROM W16.')
add_bullet('Khối 1: DELAY_I=2, DELAY_II=1, HAS_CMUL=0.')
add_paragraph('Môi trường này cho phép kiểm chứng logic trích xuất địa chỉ ROM (tw_addr), tín hiệu rot_sel, sự lan truyền của cờ valid/sync, và đo lường độ trễ toàn mạng một cách nhanh chóng. Đồng thời, thông qua việc sử dụng các test vector phi suy biến (non-degenerate), bài test sandbox đã giúp xác nhận thứ tự đầu ra tự nhiên của kiến trúc này là theo quy luật đảo bit (bit-reversed).')

add_heading('6.8. Thiết kế top-level FFT-256', 2)
add_paragraph('Mô-đun fft256_r22sdf_top.sv ghép nối 4 khối r22sdf_block_controlled liên tiếp nhau tạo thành một đường ống tính toán (datapath) liền mạch, đáp ứng chính xác sơ đồ kiến trúc Figure 4.')
table2 = doc.add_table(rows=1, cols=6)
table2.style = 'Table Grid'
hcells = table2.rows[0].cells
hcells[0].text = 'Khối'
hcells[1].text = 'DELAY_I'
hcells[2].text = 'DELAY_II'
hcells[3].text = 'Twiddle ROM'
hcells[4].text = 'HAS_CMUL'
hcells[5].text = 'Độ Trễ Đóng Góp'

def add_r(vals):
    r = table2.add_row().cells
    for i in range(6): r[i].text = str(vals[i])

add_r(['Block 0', '128', '64', 'W256', '1', '193 chu kỳ'])
add_r(['Block 1', '32', '16', 'W64', '1', '49 chu kỳ'])
add_r(['Block 2', '8', '4', 'W16', '1', '13 chu kỳ'])
add_r(['Block 3', '2', '1', 'None', '0', '3 chu kỳ'])

add_paragraph('Lý giải độ trễ mạng (Latency): Mỗi khối đóng góp độ trễ bằng tổng bộ trễ bên trong (DELAY_I + DELAY_II). Nếu khối đó có bộ nhân phức (HAS_CMUL=1), nó sẽ tiêu tốn thêm 1 chu kỳ từ thanh ghi chốt ngõ ra của CMUL. Tổng độ trễ toàn hệ thống từ khi mẫu dữ liệu đầu tiên bước vào cho đến khi xuất hiện ở ngõ ra là 193 + 49 + 13 + 3 = 258 chu kỳ clock.')

add_heading('6.9. Kiểm chứng chức năng RTL', 2)
add_paragraph('Thiết kế đã được kiểm chứng chức năng bằng mô phỏng RTL so với mô hình tham chiếu lý tưởng (golden model) viết trên nền tảng Python/NumPy. Bảy (7) bài kiểm tra vector bao gồm mảng số không, xung impulse, tín hiệu DC không đổi, các tone đơn tần số rời rạc, và dải số phức ngẫu nhiên (random complex).')
add_paragraph('Ngoài ra, hệ thống cũng đã được vượt qua:')
add_bullet('Random Stress Test: Chạy liên tiếp 20 khung dữ liệu (frames) ngẫu nhiên.')
add_bullet('Back-to-back Test: Bơm 3 khung dữ liệu chảy liên tục không có chu kỳ nghỉ để chứng minh tính đúng đắn của logic reset cục bộ (sync_in).')
add_paragraph('Kết quả xác nhận tín hiệu ra có dạng Bit-Reversed order. Do ảnh hưởng của 8 tầng chia tỷ lệ (shift-right 1 bit ở radix-2), sự hao hụt các bit trọng số thấp (LSB) tạo ra sai số lượng tử hóa. Các phép kiểm chứng ghi nhận sai số lớn nhất luôn nhỏ hơn hoặc bằng 5 LSB, thiết lập ngưỡng dung sai (error tolerance) 5 LSB cho toàn hệ thống. Cần lưu ý, đây là hoạt động kiểm chứng chức năng (functional verification), không phải là chứng minh hình thức bằng toán học (formal verification).')

add_heading('6.10. Hướng dẫn quan sát waveform bằng GTKWave', 2)
add_paragraph('Testbench được thiết lập để xuất dữ liệu trạng thái sóng dạng vcd. Có thể khởi tạo mô phỏng bằng công cụ Icarus Verilog thông qua các lệnh biên dịch iverilog và vvp. Phần mềm GTKWave được sử dụng để trực quan hóa quá trình hoạt động.')
add_paragraph('Các tín hiệu quan trọng cần được nhóm lại trong Waveform để quan sát gồm clk, rst_n, cờ trạng thái valid/sync ở cấp độ khối (block-level), dữ liệu thực/ảo và các bit điều khiển (phase_sel_i, rot_sel). Bằng cách sử dụng thanh đo (marker), người phát triển có thể đếm chính xác khoảng cách 258 chu kỳ giữa sync_in và sync_out, hoặc quan sát sự lệch pha giữa cờ valid và dữ liệu (nếu có).')
add_paragraph('[Hình 6.3: Placeholder - Waveform kiểm tra latency 258 cycles]')

add_heading('6.11. Tổng kết chương', 2)
add_paragraph('Chương này đã làm rõ quá trình triển khai sơ đồ khối kiến trúc R2²SDF ở cấu hình 256 điểm thành các mô-đun RTL cụ thể và có khả năng tổng hợp (synthesizable). Qua quy trình xây dựng từ dưới lên và bổ sung các cờ giao thức streaming thực tế, hệ thống cấu trúc FFT-256 đã khớp hoàn toàn với thiết kế lý thuyết gốc ở Hình 4 (datapath level).')
add_paragraph('Sự cẩn trọng trong phương pháp xây dựng module điều khiển cục bộ (r22sdf_block_controlled) đã mang lại tính đồng bộ hoàn hảo trong toàn mạch. Kết quả là toàn bộ hệ thống cuối cùng đã được kiểm chứng chức năng thành công bằng mô phỏng RTL (functionally verified by RTL simulation) trước sự đối chiếu nghiêm ngặt với mô hình chuẩn Python/NumPy.')

doc.save('Chuong_6_Trien_khai_RTL_R22SDF.docx')
print("Done")
